"""Validate anonymous intake workbooks locally, with no remote references."""

from __future__ import annotations

import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import pandas as pd

from scripts.release_support import _is_identifying_key

STANDARD_TOPICS = (
    "Politics",
    "History",
    "Geography",
    "Arts",
    "Sports",
    "National Heritage",
    "Kid's World",
    "Everyday Life",
)
FIXED_TOPICS = set(STANDARD_TOPICS)
WILDCARDS = {"Wildcard 1 [rename]", "Wildcard 2 [rename]"}
COMMUNITY_COLUMNS = ("community", "language_code", "language", "country", "country_code")
REQUIRED_COLUMNS = (
    "Number",
    "Topic",
    "Riddle (Original Language)",
    "Reference Answer (Original)",
)
OPTIONAL_COLUMNS = (
    "Riddle (English Translation)",
    "Hints (Optional)",
    "Comments (Optional)",
    "Reference Answer (English)",
)


def text(value) -> str:
    return "" if pd.isna(value) else str(value).strip()


# Characters that are invisible in a spreadsheet cell but make two labels differ.
INVISIBLE = dict.fromkeys(map(ord, "\u200b\u200c\u200d\u2060\ufeff"))
# Values that spreadsheet libraries write on their own; anything else may be a person's name.
TOOL_AUTHORS = {"", "openpyxl"}


def label(value) -> str:
    """Normalize a structural label (topic, number, metadata key) for comparison.

    Riddle and answer text is left as written: scripts such as Persian and many Indic
    scripts use zero-width joiners on purpose.
    """
    return " ".join(unicodedata.normalize("NFC", text(value)).translate(INVISIBLE).split())


MAX_LISTED_ROWS = 10


def _format_rows(rows: list[int]) -> str:
    """List worksheet rows, abbreviating long lists such as an untouched template."""
    if len(rows) <= MAX_LISTED_ROWS:
        return str(rows)
    return f"{rows[:MAX_LISTED_ROWS]} and {len(rows) - MAX_LISTED_ROWS} more"


@dataclass
class ValidationResult:
    metadata: dict = field(default_factory=dict)
    frame: pd.DataFrame = field(default_factory=pd.DataFrame)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _identifying(column) -> bool:
    return _is_identifying_key(column) or any(
        token in str(column).casefold() for token in ("author", "name")
    )


def _sheet(book: pd.ExcelFile, name: str, result: ValidationResult) -> pd.DataFrame:
    raw = book.parse(name, header=None, dtype=object, keep_default_na=False)
    if raw.empty:
        result.errors.append(f"{name}: sheet is empty; keep the template's header row.")
        return pd.DataFrame()
    headers = [text(value) for value in raw.iloc[0]]
    duplicates = sorted(header for header, count in Counter(headers).items() if count > 1)
    if duplicates:
        result.errors.append(f"{name}: duplicate column headers {duplicates}; keep one of each.")
    identifying = [header for header in headers if _identifying(header)]
    if identifying:
        result.errors.append(
            f"{name}: identity/attribution columns are not accepted: {identifying}. "
            "Remove these columns and their contents."
        )
    if duplicates:
        return pd.DataFrame()
    frame = raw.iloc[1:].copy()
    frame.columns = headers
    return frame


def _personal_info(path: Path) -> list[str]:
    """Describe author names that Excel stores in the file itself, outside any cell."""
    from openpyxl import load_workbook

    found = []
    book = load_workbook(path)
    for name in ("creator", "lastModifiedBy"):
        value = text(getattr(book.properties, name, None))
        if value not in TOOL_AUTHORS:
            found.append(f"document property {name} is {value!r}")
    commented = [
        f"{sheet.title}!{cell.coordinate}"
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.comment is not None
    ]
    if commented:
        found.append(f"cell comments (which record their author) at {_format_rows(commented)}")
    with ZipFile(path) as archive:
        if "docProps/app.xml" in archive.namelist():
            extended = ElementTree.fromstring(archive.read("docProps/app.xml"))
            for element in extended:
                name = element.tag.rsplit("}", 1)[-1]
                if name in {"Company", "Manager"} and text(element.text):
                    found.append(f"document property {name} is {text(element.text)!r}")
    return found


def validate(path: Path) -> ValidationResult:
    """Return the same blocking errors and review warnings for check and import."""
    result = ValidationResult()
    if path.suffix.lower() != ".xlsx":
        result.errors.append(
            "Use an .xlsx workbook with Community and Riddles sheets; CSV cannot "
            "carry the required anonymous community metadata."
        )
        return result
    try:
        with pd.ExcelFile(path, engine="openpyxl") as book:
            missing = sorted({"Community", "Riddles"} - set(book.sheet_names))
            if missing:
                result.errors.append(
                    f"Workbook is missing sheets: {missing}. Start from the anonymous template."
                )
                return result
            frames = {name: _sheet(book, name, result) for name in book.sheet_names}
    except (ValueError, BadZipFile, OSError, KeyError, ParseError) as error:
        result.errors.append(
            f"Cannot read workbook {path}: {error}. Save it as a valid .xlsx file."
        )
        return result

    personal = _personal_info(path)
    if personal:
        result.errors.append(
            "Workbook stores personal information outside the cells: "
            + "; ".join(personal)
            + ". A pull request makes it public. Remove it with "
            "`uv run python -m submissions clean YOUR_WORKBOOK.xlsx`, then check again."
        )

    metadata = frames["Community"]
    metadata = metadata.loc[metadata.apply(lambda column: column.map(text)).ne("").any(axis=1)]
    if len(metadata) != 1:
        result.errors.append(
            f"Community: expected exactly one metadata row, found {len(metadata)}."
        )
    else:
        result.metadata = {key: text(value) for key, value in metadata.iloc[0].items()}
        for key in ("community", "language_code"):
            if not result.metadata.get(key):
                result.errors.append(f"Community: {key} is required; fill in its metadata cell.")
        community = result.metadata.get("community", "")
        if community and not re.fullmatch(r"[a-z0-9]+(?:_[a-z0-9]+)*", community):
            result.errors.append(
                "Community: community must use lowercase ASCII letters, digits, and "
                "single underscores (for example, example_country)."
            )
        for key in ("language", "country", "country_code"):
            if not result.metadata.get(key):
                result.warnings.append(f"Community: fill in {key} for maintainer review.")
        code = result.metadata.get("language_code", "")
        if code and not re.fullmatch(r"[a-z]{3}", code):
            result.warnings.append(
                f"Community: language_code {code!r} is not a three-letter lowercase ISO 639-3 "
                "code (for example, tam for Tamil)."
            )
        country = result.metadata.get("country_code", "")
        if country and not re.fullmatch(r"[A-Z]{3}", country):
            result.warnings.append(
                f"Community: country_code {country!r} is not a three-letter uppercase "
                "ISO 3166-1 code (for example, IND for India)."
            )
        extra = sorted(set(metadata.columns) - set(COMMUNITY_COLUMNS))
        if extra:
            result.warnings.append(f"Community: extra columns are not imported: {extra}.")

    frame = frames["Riddles"]
    # Formatted but empty rows below the data are not riddles.
    if not frame.empty:
        frame = frame.loc[frame.apply(lambda column: column.map(text)).ne("").any(axis=1)]
    if "Topic" in frame:
        frame = frame.assign(Topic=frame["Topic"].map(label))
        canonical = {topic.casefold(): topic for topic in STANDARD_TOPICS}
        recased = sorted(
            {
                topic
                for topic in frame["Topic"]
                if topic.casefold() in canonical and topic not in FIXED_TOPICS
            }
        )
        if recased:
            frame = frame.assign(
                Topic=frame["Topic"].map(lambda topic: canonical.get(topic.casefold(), topic))
            )
            result.warnings.append(
                f"Riddles: standard topics {recased} were matched to their template spelling."
            )
    result.frame = frame
    missing = sorted(set(REQUIRED_COLUMNS) - set(frame.columns))
    if missing:
        result.errors.append(f"Riddles: missing required columns: {', '.join(missing)}.")
    optional_missing = sorted(set(OPTIONAL_COLUMNS) - set(frame.columns))
    if optional_missing:
        result.warnings.append(f"Riddles: optional columns are absent: {optional_missing}.")
    extra = sorted(set(frame.columns) - set(REQUIRED_COLUMNS) - set(OPTIONAL_COLUMNS))
    if extra:
        result.warnings.append(f"Riddles: extra columns are not imported: {extra}.")
    if len(frame) != 100:
        result.errors.append(f"Riddles: expected 100 riddles, found {len(frame)}.")
    for column in REQUIRED_COLUMNS:
        if column not in frame:
            continue
        blank = frame[column].map(text).eq("")
        if blank.any():
            rows = (frame.index[blank] + 1).tolist()
            result.errors.append(
                f"Riddles: {column} is blank at worksheet rows {_format_rows(rows)}."
            )

    numbers = []
    if "Number" in frame:
        for index, value in frame["Number"].items():
            value = label(value)
            if re.fullmatch(r"[0-9]+(?:\.0+)?", value):
                numbers.append(int(value.split(".")[0]))
            else:
                numbers.append(None)
                if value:
                    result.errors.append(
                        f"Riddles: worksheet row {index + 1}: Number must be an integer "
                        f"from 1 through 100; found {value!r}."
                    )
        counts = Counter(number for number in numbers if number is not None)
        duplicates = sorted(number for number, count in counts.items() if count > 1)
        if duplicates:
            result.errors.append(
                f"Riddles: Number values must be unique; duplicates: {duplicates}."
            )
        missing_numbers = sorted(set(range(1, 101)) - set(counts))
        unexpected = sorted(set(counts) - set(range(1, 101)))
        if missing_numbers or unexpected:
            result.errors.append(
                "Riddles: Number values must cover 1 through 100 exactly once; "
                f"missing: {missing_numbers}; out of range: {unexpected}."
            )
        elif numbers != list(range(1, 101)):
            result.warnings.append(
                "Riddles: rows are not ordered by Number; IDs use Number, not worksheet position."
            )

    if "Topic" in frame:
        topics = frame["Topic"]
        counts = topics.value_counts().to_dict()
        bad_counts = {topic: count for topic, count in counts.items() if count != 10}
        if bad_counts:
            result.errors.append(
                f"Riddles: each topic must contain exactly 10 riddles; found {bad_counts}."
            )
        present = set(topics) - {""}
        missing_topics = sorted(FIXED_TOPICS - present)
        if missing_topics:
            result.errors.append(
                f"Riddles: required standard topics are missing: {missing_topics}."
            )
        replacements = present - FIXED_TOPICS - WILDCARDS
        if len(replacements) != 2:
            result.errors.append(
                f"Riddles: include exactly two distinct renamed wildcard topics; found {sorted(replacements)}."
            )
        if present & WILDCARDS:
            result.errors.append(
                "Riddles: rename both wildcard topics before submission; placeholders remain: "
                f"{sorted(present & WILDCARDS)}."
            )
        moved = [
            number
            for number, topic in zip(numbers, topics)
            if number is not None
            and 1 <= number <= 80
            and topic != STANDARD_TOPICS[(number - 1) // 10]
        ]
        if moved:
            result.warnings.append(
                f"Riddles: standard topics differ from the template's Number blocks at {moved}. "
                "Review the assignments; existing IDs are retained by Number."
            )

    for column in ("Reference Answer (Original)", "Reference Answer (English)"):
        if column not in frame:
            continue
        groups = defaultdict(list)
        for index, value in frame[column].items():
            normalized = " ".join(text(value).casefold().split())
            if normalized:
                groups[normalized].append(index + 1)
        duplicates = [rows for rows in groups.values() if len(rows) > 1]
        if duplicates:
            result.warnings.append(
                f"Riddles: {column} has duplicate answers (case and whitespace ignored) "
                f"at worksheet row groups {duplicates}. Review them; repeated answers are allowed."
            )
    clue, answer = "Riddle (Original Language)", "Reference Answer (Original)"
    if clue in frame and answer in frame:
        revealing = [
            index + 1
            for index, row in frame.iterrows()
            if len(label(row[answer])) >= 3
            and label(row[answer]).casefold() in label(row[clue]).casefold()
        ]
        if revealing:
            result.warnings.append(
                f"Riddles: the reference answer appears in its own riddle at worksheet rows "
                f"{_format_rows(revealing)}. Review them; a riddle should not name its answer."
            )
        groups = defaultdict(list)
        for index, value in frame[clue].items():
            normalized = label(value).casefold()
            if normalized:
                groups[normalized].append(index + 1)
        repeated = [rows for rows in groups.values() if len(rows) > 1]
        if repeated:
            result.warnings.append(
                f"Riddles: the same riddle text appears at worksheet row groups {repeated}. "
                "Review them for copying mistakes."
            )
    return result
