"""Validate and maintain the canonical, anonymous riddle corpus."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

import pandas as pd

from . import evaluation_fingerprint
from .validation import (
    COMMUNITY_COLUMNS,
    FIXED_TOPICS,
    STANDARD_TOPICS,
    WILDCARDS,
    _personal_info,
    label,
    validate,
)
from .validation import text as _text

ROOT = Path(__file__).resolve().parents[1]
CANONICAL = ROOT / "data/canonical/riddles.jsonl"
METADATA = ROOT / "data/canonical/corpus_metadata.csv"
LEDGER = ROOT / "data/corrections.csv"
COMMUNITIES_DOC = ROOT / "docs/communities.md"
FIELDS = [
    "riddle_id",
    "community",
    "language",
    "topic",
    "is_wildcard",
    "riddle_original",
    "riddle_english",
    "riddle_hint",
    "riddle_note",
    "reference_answer_original",
    "reference_answer_english",
    "is_benchmarked",
]
EDITABLE = set(FIELDS) - {
    "riddle_id",
    "community",
    "language",
    "topic",
    "is_wildcard",
    "is_benchmarked",
}
PUBLIC_FIELDS = [field for field in FIELDS if field != "is_benchmarked"]
LEDGER_FIELDS = ["riddle_id", "action", "field", "old_value", "new_value", "reason"]


def load_rows(path: Path | None = None) -> list[dict]:
    if path is None:
        path = CANONICAL
    with path.open(encoding="utf-8") as stream:
        return [json.loads(line) for line in stream if line.strip()]


def fingerprint(rows: list[dict]) -> str:
    canonical = "".join(
        json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n"
        for row in rows
    )
    return hashlib.sha256(canonical.encode()).hexdigest()


def check(args) -> int:
    result = validate(args.input)
    community = result.metadata.get("community", "")
    if not result.errors and community:
        existing = sum(row["community"] == community for row in load_rows())
        if existing:
            result.warnings.append(
                f"Community: {community} already exists with {existing} riddles; importing this "
                "workbook replaces all of them. Use a new key unless you are updating that "
                "community's set with its maintainers."
            )
    print(
        f"{'FAIL' if result.errors else 'PASS'}: {args.input} "
        f"({len(result.errors)} issue(s), {len(result.warnings)} warning(s))"
    )
    for error in result.errors:
        print(f"- {error}")
    for warning in result.warnings:
        print(f"WARNING: {warning}")
    return int(bool(result.errors))


def _slug(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", re.sub(r"['’]", "", ascii_value.lower())).strip("_")


def _new_rows(meta: dict, table: pd.DataFrame, old: list[dict]) -> list[dict]:
    community = meta["community"]
    language = meta["language_code"]
    previous = [row for row in old if row["community"] == community]
    old_by_number = {row["riddle_id"].rsplit("_", 1)[-1]: row for row in previous}
    if len(old_by_number) != len(previous):
        raise ValueError("Existing community has ambiguous riddle numbers.")
    numbered = [
        (int(re.sub(r"\.0+$", "", label(n))), label(t))
        for n, t in zip(table["Number"], table["Topic"])
    ]
    wildcard_order = [t for _, t in sorted(numbered) if t not in FIXED_TOPICS]
    wildcard_topics = list(dict.fromkeys(wildcard_order))
    output = []
    for _, s in table.iterrows():
        topic = label(s.get("Topic"))
        number = label(s.get("Number"))
        number_slug = str(int(re.sub(r"\.0+$", "", number))).zfill(3)
        prior = old_by_number.get(number_slug)
        # Topics in scripts without an ASCII spelling fall back to their wildcard position.
        topic_slug = _slug(topic) or f"wildcard_{wildcard_topics.index(topic) + 1}"
        rid = prior["riddle_id"] if prior else f"{community}_{topic_slug}_{number_slug}"
        row = {
            "riddle_id": rid,
            "community": community,
            "language": language,
            "topic": topic,
            "is_wildcard": topic in wildcard_topics,
            "riddle_original": _text(s.get("Riddle (Original Language)")),
            "riddle_english": _text(s.get("Riddle (English Translation)")) or None,
            "riddle_hint": _text(s.get("Hints (Optional)")) or None,
            "riddle_note": _text(s.get("Comments (Optional)")) or None,
            "reference_answer_original": _text(s.get("Reference Answer (Original)")),
            "reference_answer_english": _text(s.get("Reference Answer (English)")) or None,
            "is_benchmarked": False,
        }
        row["is_benchmarked"] = bool(
            prior
            and prior.get("is_benchmarked", False)
            and evaluation_fingerprint(prior) == evaluation_fingerprint(row)
        )
        output.append(row)
    if len({r["riddle_id"] for r in output}) != len(output):
        raise ValueError("Rows produce duplicate stable riddle IDs.")
    return output


def import_submission(args) -> int:
    result = validate(args.input)
    if result.errors:
        raise ValueError("Submission fails validation: " + "; ".join(result.errors))
    for warning in result.warnings:
        print(f"WARNING: {warning}")
    meta, frame = result.metadata, result.frame
    current = load_rows()
    incoming = _new_rows(meta, frame, current)
    comm = meta["community"]
    old = [r for r in current if r["community"] == comm]
    other = [r for r in current if r["community"] != comm]
    ids = {r["riddle_id"] for r in current}
    newids = {r["riddle_id"] for r in incoming}
    if not old and ids & newids:
        raise ValueError("Incoming riddle IDs collide with another community.")
    added = len(newids - {r["riddle_id"] for r in old})
    removed = len({r["riddle_id"] for r in old} - newids)
    changed = sum(
        1
        for r in incoming
        if r["riddle_id"] in {x["riddle_id"] for x in old} and not r["is_benchmarked"]
    )
    print(
        f"Community {comm}: {len(old)} -> {len(incoming)} rows; {added} added, {removed} removed, {changed} text changes invalidate prior benchmark attachment."
    )
    if not args.write:
        print("Dry run; canonical data was not changed.")
        return 0
    if not args.approved:
        raise ValueError("Import requires --approved after maintainer review.")
    result = other + incoming
    CANONICAL.write_text(
        "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in result), encoding="utf-8"
    )
    community_metadata = pd.read_csv(METADATA).set_index("community", drop=False)
    values = {
        "language_code": _text(meta.get("language_code", "")),
        "language": _text(meta.get("language", "")),
        "country": _text(meta.get("country", "")),
        "country_code": _text(meta.get("country_code", "")),
    }
    if comm in community_metadata.index:
        for key, value in values.items():
            if value:
                community_metadata.loc[comm, key] = value
    else:
        new_meta = {column: "" for column in community_metadata.columns}
        new_meta.update({key: value for key, value in values.items() if key in new_meta})
        new_meta["community"] = comm
        community_metadata.loc[comm] = new_meta
    community_metadata.reset_index(drop=True).to_csv(METADATA, index=False)
    print(f"Wrote {CANONICAL} (fingerprint {fingerprint(result)}).")
    print(f"Updated {METADATA} for {comm}.")
    return 0


def corrections(args) -> int:
    if args.ledger.resolve() == LEDGER.resolve():
        raise ValueError(
            "data/corrections.csv is applied history; --ledger must name a separate pending review file."
        )
    current = load_rows()
    with args.ledger.open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != LEDGER_FIELDS:
            raise ValueError(f"Ledger columns must be {LEDGER_FIELDS}.")
        entries = list(reader)
    result = [dict(r) for r in current]
    index = {r["riddle_id"]: r for r in result}
    remove = set()
    errors = []
    for n, c in enumerate(entries, 2):
        rid, action, field, old, new, reason = (c[k] or "" for k in LEDGER_FIELDS)
        if not reason:
            errors.append(f"line {n}: reason required")
            continue
        if action not in {"remove", "replace", "append"}:
            errors.append(f"line {n}: invalid action {action!r}")
            continue
        row = index.get(rid)
        if action == "remove":
            if field or old or new:
                errors.append(f"line {n}: remove must not set field/values")
            elif row is None:
                errors.append(f"line {n}: cannot verify removal; riddle is absent")
            else:
                remove.add(rid)
            continue
        if field not in EDITABLE:
            errors.append(f"line {n}: non-editable field {field!r}")
            continue
        if row is None:
            errors.append(f"line {n}: riddle not found: {rid}")
            continue
        value = row.get(field) or ""
        if action == "replace":
            if not old or not new:
                errors.append(f"line {n}: replace requires old_value and new_value")
            elif value == old:
                row[field] = new
                row["is_benchmarked"] = False
            elif value != new:
                errors.append(f"line {n}: current value does not equal old_value or new_value")
        else:
            if not new:
                errors.append(f"line {n}: append requires new_value")
            elif value == new or value.endswith("\n" + new):
                pass
            elif value == old:
                row[field] = (value + "\n" if value else "") + new
                row["is_benchmarked"] = False
            else:
                errors.append(
                    f"line {n}: current value does not match old_value (append is ambiguous)"
                )
    if errors:
        raise ValueError("Corrections rejected; nothing written:\n" + "\n".join(errors))
    result = [r for r in result if r["riddle_id"] not in remove]
    print(f"Corrections: {len(current)} -> {len(result)} rows; fingerprint {fingerprint(result)}")
    if not args.write:
        print("Dry run; canonical data was not changed.")
        return 0
    if not args.approved:
        raise ValueError("Applying corrections requires --approved after maintainer review.")
    history = []
    if LEDGER.exists():
        with LEDGER.open(encoding="utf-8", newline="") as stream:
            history_reader = csv.DictReader(stream)
            if history_reader.fieldnames != LEDGER_FIELDS:
                raise ValueError("Applied correction history has invalid columns; nothing written.")
            history = list(history_reader)
    known = {tuple(entry[field] for field in LEDGER_FIELDS[:-1]) for entry in history}
    for entry in entries:
        key = tuple(entry[field] for field in LEDGER_FIELDS[:-1])
        if key not in known:
            history.append(entry)
            known.add(key)
    CANONICAL.write_text(
        "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in result), encoding="utf-8"
    )
    LEDGER.parent.mkdir(parents=True, exist_ok=True)
    with LEDGER.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=LEDGER_FIELDS)
        writer.writeheader()
        writer.writerows(history)
    return 0


def data_output(path: Path) -> Path:
    """Resolve artifacts beneath data, including through existing symlinks."""
    resolved = path.resolve()
    data_root = (ROOT / "data").resolve()
    if not resolved.is_relative_to(data_root) or resolved == data_root:
        raise ValueError(
            "Output artifacts must remain in a dedicated path beneath repository data/."
        )
    return resolved


def snapshot_audit(rows: list[dict], baseline_path: Path) -> dict:
    baseline = {
        item["riddle_id"]: item["text_sha256"]
        for line in baseline_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
        for item in [json.loads(line)]
    }
    current = {row["riddle_id"]: row for row in rows}
    current_benchmarked = {rid for rid, row in current.items() if row.get("is_benchmarked", False)}
    matching = set()
    changed = set()
    for rid in baseline.keys() & current.keys():
        if evaluation_fingerprint(current[rid]) == baseline[rid]:
            matching.add(rid)
        else:
            changed.add(rid)
    missing = baseline.keys() - current.keys()
    never_benchmarked = {
        rid
        for rid, row in current.items()
        if not row.get("is_benchmarked", False) and rid not in baseline
    }
    return {
        "original_benchmarked_riddle_ids": sorted(baseline),
        "original_benchmarked_riddle_count": len(baseline),
        "selected_benchmark_riddle_ids": sorted(current_benchmarked),
        "selected_benchmark_count": len(current_benchmarked),
        "matching_evaluated_riddle_ids": sorted(matching & current_benchmarked),
        "changed_evaluated_riddle_ids": sorted(changed),
        "removed_evaluated_riddle_ids": sorted(missing),
        "never_benchmarked_riddle_ids": sorted(never_benchmarked),
        "added_riddle_ids": sorted(current.keys() - baseline.keys()),
        "snapshot_matches": current_benchmarked == set(baseline) and not changed and not missing,
    }


def export(args) -> int:
    import shutil

    target = data_output(args.output)
    if not args.write:
        print(
            f"Dry run: export anonymous canonical riddles, manifest, card, and LICENSE to {target}."
        )
        print("No files written. Pass --write to export.")
        return 0
    license_path = ROOT / "data/LICENSE"
    if not license_path.is_file():
        raise FileNotFoundError(f"Dataset terms are required for export: {license_path}")
    if target.exists() and (not target.is_dir() or any(target.iterdir())):
        raise ValueError(
            "Export output must be an empty directory; choose a fresh path to avoid stale upload files."
        )
    rows = load_rows()
    manifest = {
        "rows": len(rows),
        "dataset_sha256": fingerprint(rows),
        "evaluation_snapshot_sha256": hashlib.sha256(args.snapshot.read_bytes()).hexdigest(),
        "evaluation_baseline_sha256": hashlib.sha256(args.baseline.read_bytes()).hexdigest(),
        **snapshot_audit(rows, args.baseline),
    }
    target.mkdir(parents=True, exist_ok=True)
    frame = pd.DataFrame([{field: row.get(field) for field in PUBLIC_FIELDS} for row in rows])
    frame.to_parquet(target / "riddles.parquet", index=False)
    (target / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    card = (
        "---\nlicense: other\nlicense_name: culturalriddles-evaluation-only\n"
        "license_link: LICENSE\nconfigs:\n- config_name: default\n"
        "  data_files:\n  - split: test\n    path: riddles.parquet\n---\n\n"
        "# Cultural Riddles\n\n"
        f"{len(rows)} anonymous riddles. Preserve the attribution and evaluation-only terms in LICENSE.\n\n"
        f"Dataset fingerprint: `{manifest['dataset_sha256']}`. "
        f"Frozen evaluation content match: `{manifest['snapshot_matches']}`.\n"
    )
    (target / "README.md").write_text(card, encoding="utf-8")
    shutil.copyfile(license_path, target / "LICENSE")
    print(f"Exported {len(rows)} rows to {target}; fingerprint {manifest['dataset_sha256']}")
    return 0


def publish(args) -> int:
    bundle = data_output(args.bundle)
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*/[A-Za-z0-9][A-Za-z0-9_.-]*", args.repo_id):
        raise ValueError("--repo-id must have the form owner/dataset-name.")
    if not args.publish:
        print(f"Dry run: upload the existing local bundle {bundle} to dataset {args.repo_id}.")
        print("No files written or external services contacted. Pass --publish to publish.")
        return 0
    if not args.native_review_approved:
        raise ValueError("Publishing requires recorded native-speaker review approval.")
    required = {"riddles.parquet", "manifest.json", "README.md", "LICENSE"}
    missing = sorted(name for name in required if not (bundle / name).is_file())
    if missing:
        raise FileNotFoundError(f"Release bundle is incomplete at {bundle}: missing {missing}")
    from huggingface_hub import HfApi

    api = HfApi()
    api.create_repo(repo_id=args.repo_id, repo_type="dataset", private=True, exist_ok=True)
    api.upload_folder(
        repo_id=args.repo_id,
        repo_type="dataset",
        folder_path=str(bundle),
        allow_patterns=sorted(required),
    )
    print(
        f"Published the reviewed dataset bundle to {args.repo_id}. Existing repository visibility is unchanged."
    )
    return 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(prog="python -m submissions")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("check")
    p.add_argument("input", type=Path)
    p.set_defaults(func=check)
    p = sub.add_parser("import")
    p.add_argument("input", type=Path)
    p.add_argument("--write", action="store_true")
    p.add_argument("--approved", action="store_true")
    p.set_defaults(func=import_submission)
    p = sub.add_parser("corrections")
    p.add_argument(
        "--ledger",
        type=Path,
        required=True,
        help="Pending reviewed corrections, separate from applied data/corrections.csv",
    )
    p.add_argument("--write", action="store_true")
    p.add_argument("--approved", action="store_true")
    p.set_defaults(func=corrections)
    for command, handler in (("export", export), ("publish", publish)):
        p = sub.add_parser(command)
        if command == "export":
            p.add_argument("--output", type=Path, default=ROOT / "data/derived/submissions/export")
            p.add_argument(
                "--snapshot", type=Path, default=ROOT / "data/canonical/benchmark_results.parquet"
            )
            p.add_argument(
                "--baseline", type=Path, default=ROOT / "data/canonical/evaluation_corpus.jsonl"
            )
            p.add_argument("--write", action="store_true")
        else:
            p.add_argument("--bundle", type=Path, default=ROOT / "data/derived/submissions/export")
            p.add_argument("--publish", action="store_true")
            p.add_argument("--repo-id", required=True)
            p.add_argument("--native-review-approved", action="store_true")
        p.set_defaults(func=handler)
    p = sub.add_parser("clean", help="Remove author names stored outside the cells")
    p.add_argument("input", type=Path)
    p.set_defaults(func=clean)
    p = sub.add_parser("communities", help="List the communities already in the corpus")
    mode = p.add_mutually_exclusive_group()
    mode.add_argument("--write", action="store_true", help="Update docs/communities.md")
    mode.add_argument("--check", action="store_true", help="Fail if docs/communities.md is stale")
    p.set_defaults(func=communities)
    p = sub.add_parser("template")
    p.add_argument(
        "--output", type=Path, default=ROOT / "data/submissions/submission_template.xlsx"
    )
    p.set_defaults(func=template)
    args = parser.parse_args(argv)
    try:
        return args.func(args)
    except (ValueError, FileNotFoundError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 2


def clean(args) -> int:
    """Remove author names that Excel stores outside the cells, in place.

    Saving through openpyxl also rewrites the extended properties (company and manager).
    """
    from openpyxl import load_workbook

    path = args.input
    if path.suffix.lower() != ".xlsx":
        raise ValueError("clean works on .xlsx workbooks.")
    found = _personal_info(path)
    book = load_workbook(path)
    for name in ("creator", "lastModifiedBy"):
        setattr(book.properties, name, "")
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.comment = None
    book.save(path)
    remaining = _personal_info(path)
    if remaining:
        raise ValueError(f"Could not remove: {'; '.join(remaining)}.")
    print(f"Cleaned {path}: removed {'; '.join(found) if found else 'nothing'}.")
    print("Cell text is unchanged; review it yourself for names before submitting.")
    return 0


def communities_markdown(rows: list[dict], metadata: pd.DataFrame) -> str:
    """Render the list of communities already in the corpus."""
    counts = Counter(row["community"] for row in rows)
    meta = metadata.set_index("community")
    lines = [
        "# Communities in Cultural Riddles",
        "",
        "Generated from `data/canonical/` by `uv run python -m submissions communities --write`; "
        "do not edit by hand.",
        "",
        "Check this list before you start. To add a community that is not here, open an "
        "[Add my community](https://github.com/Cohere-Labs-Community/CulturalRiddles/issues/new?template=add-community.yml) "
        "issue so that others can see it is in progress. To improve a community that is "
        "already here, open an issue first instead of submitting a replacement.",
        "",
        "| Community key | Language | Country | Riddles |",
        "| --- | --- | --- | ---: |",
    ]
    for community in sorted(counts):
        row = meta.loc[community] if community in meta.index else {}
        language = _text(row.get("language", "")) if len(row) else ""
        country = _text(row.get("country", "")) if len(row) else ""
        lines.append(f"| `{community}` | {language} | {country} | {counts[community]} |")
    return "\n".join(lines) + "\n"


def communities(args) -> int:
    rendered = communities_markdown(load_rows(), pd.read_csv(METADATA))
    if args.write:
        COMMUNITIES_DOC.write_text(rendered, encoding="utf-8")
        print(f"Wrote {COMMUNITIES_DOC}")
        return 0
    current = COMMUNITIES_DOC.read_text(encoding="utf-8") if COMMUNITIES_DOC.exists() else ""
    if args.check:
        if current != rendered:
            raise ValueError(
                f"{COMMUNITIES_DOC} is out of date; run `uv run python -m submissions communities --write`."
            )
        print(f"{COMMUNITIES_DOC} is up to date.")
        return 0
    print(rendered, end="")
    return 0


def template(args) -> int:
    from openpyxl import Workbook

    path = data_output(args.output)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    instructions = wb.active
    instructions.title = "Instructions"
    instructions.append(["Cultural Riddles contribution workbook"])
    instructions.append(
        [
            "Fill in community and language_code; use a lowercase community key such as example_country."
        ]
    )
    instructions.append(
        ["Do not include personal names, contributor identities, or per-riddle attribution."]
    )
    instructions.append(
        [
            "Enter exactly 100 riddles. Keep Number values 1 through 100; each topic must have 10 rows."
        ]
    )
    instructions.append(
        [
            "Keep all eight standard topics and rename the two wildcard topics to distinct community topics."
        ]
    )
    instructions.append(
        [
            "Complete both original-language text and reference answer. English fields, hints, and comments are optional."
        ]
    )
    community = wb.create_sheet("Community")
    community.append(list(COMMUNITY_COLUMNS))
    community.append(["", "", "", "", ""])
    sheet = wb.create_sheet("Riddles")
    headers = [
        "Number",
        "Topic",
        "Riddle (Original Language)",
        "Riddle (English Translation)",
        "Hints (Optional)",
        "Comments (Optional)",
        "Reference Answer (Original)",
        "Reference Answer (English)",
    ]
    sheet.append(headers)
    topics = [*STANDARD_TOPICS, *sorted(WILDCARDS)]
    for number in range(1, 101):
        sheet.append([number, topics[(number - 1) // 10], None, None, None, None, None, None])
    wb.save(path)
    print(f"Wrote anonymous template {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
