from argparse import Namespace

import pandas as pd
import pytest
from openpyxl import load_workbook

from submissions import __main__ as app
from submissions.validation import OPTIONAL_COLUMNS, STANDARD_TOPICS, validate


@pytest.fixture
def intake(tmp_path):
    path = tmp_path / "anonymous.xlsx"
    metadata = pd.DataFrame(
        [
            {
                "community": "example_country",
                "language_code": "eng",
                "language": "English",
                "country": "Example country",
                "country_code": "XXX",
            }
        ]
    )
    topics = [*STANDARD_TOPICS, "Food", "Traditions"]
    rows = pd.DataFrame(
        [
            {
                "Number": number,
                "Topic": topics[(number - 1) // 10],
                "Riddle (Original Language)": f"Clue {number}",
                "Reference Answer (Original)": f"Answer {number}",
                **{column: "" for column in OPTIONAL_COLUMNS},
            }
            for number in range(1, 101)
        ]
    )

    def save():
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            metadata.to_excel(writer, sheet_name="Community", index=False)
            rows.to_excel(writer, sheet_name="Riddles", index=False)
        return path

    rows["Number"] = rows["Number"].astype(object)
    return path, metadata, rows, save


def test_valid_intake_check_import_dry_run_and_approval_guard(intake, monkeypatch, capsys):
    path, _, _, save = intake
    save()
    canonical = path.parent / "riddles.jsonl"
    metadata = path.parent / "metadata.csv"
    canonical.write_text("unchanged corpus\n")
    metadata.write_text("unchanged metadata\n")
    monkeypatch.setattr(app, "CANONICAL", canonical)
    monkeypatch.setattr(app, "METADATA", metadata)
    monkeypatch.setattr(app, "load_rows", lambda: [])
    assert validate(path).errors == []
    assert validate(path).warnings == []
    assert app.main(["check", str(path)]) == 0
    assert app.main(["import", str(path)]) == 0
    assert "Dry run; canonical data was not changed." in capsys.readouterr().out
    assert app.main(["import", str(path), "--write"]) == 2
    assert "requires --approved" in capsys.readouterr().err
    assert canonical.read_text() == "unchanged corpus\n"
    assert metadata.read_text() == "unchanged metadata\n"


@pytest.mark.parametrize(
    "column,value,message",
    [
        ("Number", 1.5, "Number must be an integer"),
        ("Number", 101, "out of range: [101]"),
        ("Number", "002", "duplicates: [2]"),
        ("Number", " ", "Number is blank at worksheet rows [2]"),
        ("Topic", " ", "Topic is blank at worksheet rows [2]"),
        ("Topic", "Other", "each topic must contain exactly 10"),
        ("Riddle (Original Language)", " ", "Riddle (Original Language) is blank"),
        ("Reference Answer (Original)", "", "Reference Answer (Original) is blank"),
    ],
)
def test_malformed_rows_fail_check_and_approved_import_before_corpus_read(
    intake, monkeypatch, capsys, column, value, message
):
    path, _, rows, save = intake
    rows.loc[0, column] = value
    save()

    def forbidden_read():
        pytest.fail("Invalid intake must fail before reading or writing canonical data")

    monkeypatch.setattr(app, "load_rows", forbidden_read)
    assert any(message in error for error in validate(path).errors)
    assert app.main(["check", str(path)]) == 1
    assert message in capsys.readouterr().out
    assert app.main(["import", str(path), "--approved", "--write"]) == 2
    assert message in capsys.readouterr().err


def test_many_blank_rows_are_abbreviated(intake):
    path, _, rows, save = intake
    rows["Riddle (Original Language)"] = ""
    save()

    errors = [e for e in validate(path).errors if "Riddle (Original Language) is blank" in e]
    assert errors == [
        "Riddles: Riddle (Original Language) is blank at worksheet rows "
        "[2, 3, 4, 5, 6, 7, 8, 9, 10, 11] and 90 more."
    ]


@pytest.mark.parametrize(
    "column", ["Number", "Topic", "Riddle (Original Language)", "Reference Answer (Original)"]
)
def test_required_columns_are_blocking(intake, column):
    path, _, rows, save = intake
    rows.drop(columns=column, inplace=True)
    save()
    assert any(
        "missing required columns" in error and column in error for error in validate(path).errors
    )


@pytest.mark.parametrize(
    "key,value", [("community", ""), ("community", "Private Name"), ("language_code", "")]
)
def test_metadata_needed_by_import_also_fails_check(intake, key, value):
    path, metadata, _, save = intake
    metadata.loc[0, key] = value
    save()
    assert any(key in error for error in validate(path).errors)


def test_exact_row_count_and_complete_number_range_are_required(intake):
    path, _, rows, save = intake
    rows.drop(index=99, inplace=True)
    save()
    errors = validate(path).errors
    assert any("expected 100 riddles, found 99" in error for error in errors)
    assert any("missing: [100]" in error for error in errors)


@pytest.mark.parametrize("replacement", ["Wildcard 1 [rename]", "Food", "Politics"])
def test_two_distinct_renamed_wildcards_are_required(intake, replacement):
    path, _, rows, save = intake
    rows.loc[rows["Topic"].eq("Traditions"), "Topic"] = replacement
    save()
    assert any(
        "exactly two distinct renamed wildcard topics" in error for error in validate(path).errors
    )


def test_missing_standard_topic_is_blocking(intake):
    path, _, rows, save = intake
    rows.loc[rows["Topic"].eq("Politics"), "Topic"] = "Other"
    save()
    assert any(
        "required standard topics are missing: ['Politics']" in error
        for error in validate(path).errors
    )


def test_legacy_review_warnings_do_not_block_intake(intake, capsys):
    path, _, rows, save = intake
    rows.loc[0, "Reference Answer (Original)"] = "  ANSWER   2  "
    rows.loc[0, "Topic"], rows.loc[10, "Topic"] = rows.loc[10, "Topic"], rows.loc[0, "Topic"]
    rows["Extra context"] = "not imported"
    rows.drop(columns="Hints (Optional)", inplace=True)
    save()
    result = validate(path)
    assert not result.errors
    for message in ("duplicate answers", "Number blocks", "extra columns", "optional columns"):
        assert any(message in warning for warning in result.warnings)
    assert app.check(Namespace(input=path)) == 0
    assert "WARNING:" in capsys.readouterr().out


def test_literal_na_is_content_and_surrounding_topic_spaces_are_normalized(intake):
    path, _, rows, save = intake
    rows.loc[0, "Reference Answer (Original)"] = "NA"
    rows["Topic"] = rows["Topic"].map(lambda topic: f"  {topic}  ")
    save()
    result = validate(path)
    assert result.errors == []
    assert result.frame.iloc[0]["Reference Answer (Original)"] == "NA"


def test_duplicate_headers_are_not_silently_renamed(intake):
    path, _, _, save = intake
    save()
    book = load_workbook(path)
    book["Riddles"].cell(1, 9, "Number")
    book.save(path)
    assert any("duplicate column headers ['Number']" in error for error in validate(path).errors)


def test_identity_columns_in_extra_sheets_are_rejected(intake):
    path, _, _, save = intake
    save()
    book = load_workbook(path)
    book.create_sheet("Private").append(["Email"])
    book.save(path)
    assert any("Private: identity/attribution" in error for error in validate(path).errors)


def test_missing_sheets_multiple_metadata_rows_and_corrupt_files_are_actionable(intake):
    path, metadata, _, save = intake
    metadata.loc[1] = metadata.loc[0]
    save()
    assert any("exactly one metadata row, found 2" in error for error in validate(path).errors)
    book = load_workbook(path)
    del book["Community"]
    book.save(path)
    assert "missing sheets" in validate(path).errors[0]
    path.write_text("not an Excel workbook")
    assert "Cannot read workbook" in validate(path).errors[0]


def test_csv_is_not_a_complete_anonymous_intake(tmp_path):
    path = tmp_path / "riddles.csv"
    path.write_text("Number,Topic\n1,Politics\n")
    assert "CSV cannot carry" in validate(path).errors[0]


def test_numeric_number_spelling_normalizes_to_the_same_id(intake):
    path, _, rows, save = intake
    rows["Number"] = rows["Number"].map(lambda number: f"{number:03}.0")
    save()
    result = validate(path)
    assert not result.errors
    incoming = app._new_rows(result.metadata, result.frame, [])
    assert incoming[0]["riddle_id"] == "example_country_politics_001"
